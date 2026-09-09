#!/usr/bin/env python3
"""Phase 1 Ingestion Pipeline Verification Script.

Usage:
    conda activate tri
    python scripts/verify_phase1.py
"""

import io
import os
import sys
import struct
import asyncio
from pathlib import Path

# Ensure UTF-8 output on Windows console
if sys.platform == "win32":
    sys.stdout.reconfigure(encoding="utf-8")

ROOT_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT_DIR / "backend"))

# Colors for terminal output
GREEN = "\033[92m"
RED = "\033[91m"
YELLOW = "\033[93m"
CYAN = "\033[96m"
BOLD = "\033[1m"
RESET = "\033[0m"


def print_step(name: str):
    print(f"\n{BOLD}{CYAN}>>> [CHECK] {name}{RESET}")


async def run_verification():
    print(f"\n{BOLD}================================================================{RESET}")
    print(f"{BOLD}   SIH26155 GenAI Platform — Phase 1 Ingestion Verification     {RESET}")
    print(f"{BOLD}================================================================{RESET}")

    from httpx import ASGITransport, AsyncClient
    from PIL import Image
    from pypdf import PdfWriter
    import docx
    import pptx

    from app.main import app
    from app.core.config import get_settings
    from app.core.security import decrypt_bytes, compute_sha256
    from app.core.database import get_db
    from app.models.audit_log import Base
    from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

    # Set up in-memory test database and upload directory
    test_engine = create_async_engine("sqlite+aiosqlite:///:memory:", echo=False)
    TestSessionLocal = async_sessionmaker(bind=test_engine, class_=AsyncSession, expire_on_commit=False)

    test_upload_dir = ROOT_DIR / "tmp_verification_uploads"
    test_upload_dir.mkdir(parents=True, exist_ok=True)
    get_settings().UPLOAD_STORAGE_PATH = str(test_upload_dir)

    async with test_engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)

    async def override_get_db():
        async with TestSessionLocal() as s:
            try:
                yield s
            finally:
                await s.close()

    app.dependency_overrides[get_db] = override_get_db

    all_passed = True

    try:
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            # 1. Plain Text / Markdown Ingestion
            print_step("1. Ingest Plain Text / Markdown (.md)")
            text_content = b"# Tactical SITREP\n\nAll surveillance sectors operational. Egress blocked."
            res = await client.post("/api/v1/ingest/upload", files={"file": ("sitrep.md", text_content, "text/markdown")})
            assert res.status_code == 201, f"Failed text upload: {res.text}"
            doc = res.json()["document"]
            print(f"  {GREEN}[OK]{RESET} Text ingested: doc_id={doc['doc_id']}")
            print(f"       Checksum: {doc['checksum'][:16]}... | Headings: {len(doc['structural_metadata']['headings'])}")

            # 2. PDF Document Ingestion
            print_step("2. Ingest PDF Document (.pdf)")
            writer = PdfWriter()
            writer.add_blank_page(width=300, height=300)
            pdf_stream = io.BytesIO()
            writer.write(pdf_stream)
            res = await client.post("/api/v1/ingest/upload", files={"file": ("intel_doc.pdf", pdf_stream.getvalue(), "application/pdf")})
            assert res.status_code == 201, f"Failed PDF upload: {res.text}"
            doc_pdf = res.json()["document"]
            print(f"  {GREEN}[OK]{RESET} PDF ingested: doc_id={doc_pdf['doc_id']} (Pages: {doc_pdf['structural_metadata']['total_pages']})")

            # 3. DOCX Ingestion
            print_step("3. Ingest Microsoft Word (.docx)")
            doc_obj = docx.Document()
            doc_obj.add_heading("Operational Protocol", level=1)
            doc_obj.add_paragraph("Encrypted at rest with AES-256-GCM.")
            docx_stream = io.BytesIO()
            doc_obj.save(docx_stream)
            res = await client.post("/api/v1/ingest/upload", files={"file": ("protocol.docx", docx_stream.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
            assert res.status_code == 201, f"Failed DOCX upload: {res.text}"
            doc_docx = res.json()["document"]
            print(f"  {GREEN}[OK]{RESET} DOCX ingested: doc_id={doc_docx['doc_id']}")

            # 4. PPTX Ingestion
            print_step("4. Ingest Microsoft PowerPoint (.pptx)")
            prs = pptx.Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[0])
            slide.shapes.title.text = "Air-Gapped Transformation Enclave"
            pptx_stream = io.BytesIO()
            prs.save(pptx_stream)
            res = await client.post("/api/v1/ingest/upload", files={"file": ("briefing.pptx", pptx_stream.getvalue(), "application/vnd.openxmlformats-officedocument.presentationml.presentation")})
            assert res.status_code == 201, f"Failed PPTX upload: {res.text}"
            doc_pptx = res.json()["document"]
            print(f"  {GREEN}[OK]{RESET} PPTX ingested: doc_id={doc_pptx['doc_id']}")

            # 5. Image Ingestion (PNG)
            print_step("5. Ingest Image Document (.png)")
            img = Image.new("RGB", (250, 100), color=(30, 40, 60))
            img_stream = io.BytesIO()
            img.save(img_stream, format="PNG")
            res = await client.post("/api/v1/ingest/upload", files={"file": ("radar.png", img_stream.getvalue(), "image/png")})
            assert res.status_code == 201, f"Failed Image upload: {res.text}"
            doc_img = res.json()["document"]
            print(f"  {GREEN}[OK]{RESET} Image ingested: doc_id={doc_img['doc_id']} (Dims: {doc_img['structural_metadata']['image_dimensions']})")

            # 6. Audio Ingestion (WAV)
            print_step("6. Ingest Audio Media (.wav)")
            wav_raw = b"\x00\x00" * 16000
            wav_header = struct.pack(
                "<4sI4s4sIHHIIHH4sI",
                b"RIFF", 36 + len(wav_raw), b"WAVE", b"fmt ", 16, 1, 1, 16000, 32000, 2, 16, b"data", len(wav_raw)
            )
            res = await client.post("/api/v1/ingest/upload", files={"file": ("dispatch.wav", wav_header + wav_raw, "audio/wav")})
            assert res.status_code == 201, f"Failed Audio upload: {res.text}"
            doc_audio = res.json()["document"]
            print(f"  {GREEN}[OK]{RESET} Audio ingested: doc_id={doc_audio['doc_id']} (Duration: {doc_audio['structural_metadata']['duration_seconds']}s)")

            # 7. Encryption-At-Rest Verification
            print_step("7. Verify AES-256 At-Rest Disk Encryption")
            test_raw = b"RESTRICTED DEFENSE ENCLAVE CONTENT - CANNOT BE LEAKED TO DISK IN PLAINTEXT"
            res_enc = await client.post("/api/v1/ingest/upload", files={"file": ("top_secret.txt", test_raw, "text/plain")})
            enc_doc_id = res_enc.json()["document"]["doc_id"]

            v_res = await client.get(f"/api/v1/ingest/documents/{enc_doc_id}/verify-encryption")
            v_data = v_res.json()
            assert v_data["is_unreadable_ciphertext"] is True
            assert v_data["decryption_successful"] is True

            # Direct inspection of disk file
            disk_file = Path(v_data["disk_path"])
            raw_on_disk = disk_file.read_bytes()
            assert test_raw not in raw_on_disk, "CRITICAL ERROR: Plaintext found on disk!"
            decrypted = decrypt_bytes(raw_on_disk)
            assert decrypted == test_raw
            print(f"  {GREEN}[OK]{RESET} AES-256-GCM verified: Plaintext absent from disk file")
            print(f"       File size: {len(raw_on_disk)} bytes | Preview: {v_data['sample_hex_preview'][:32]}...")

            # 8. Tamper-Evident Audit Log Verification
            print_step("8. Verify Append-Only Cryptographic Audit Log")
            chain_res = await client.get("/api/v1/audit/verify-chain")
            chain_data = chain_res.json()
            assert chain_data["valid"] is True
            print(f"  {GREEN}[OK]{RESET} Audit hash chain verified: {chain_data['total_records']} records chained")
            print(f"       Latest Hash: {chain_data['latest_hash']}")

            # 9. Error Handling: Corrupted & Unsupported Formats
            print_step("9. Verify Error Handling (Corrupted & Unsupported Formats)")
            corrupt_res = await client.post("/api/v1/ingest/upload", files={"file": ("bad.pdf", b"%PDF-truncated", "application/pdf")})
            assert corrupt_res.status_code == 400
            print(f"  {GREEN}[OK]{RESET} Corrupted file rejected: HTTP 400 (Detail: {corrupt_res.json()['detail']['message']})")

            bad_ext_res = await client.post("/api/v1/ingest/upload", files={"file": ("script.sh", b"#!/bin/bash\nrm -rf /", "application/x-sh")})
            assert bad_ext_res.status_code == 415
            print(f"  {GREEN}[OK]{RESET} Unsupported format rejected: HTTP 415 (Detail: {bad_ext_res.json()['detail']['content_type']})")

    except Exception as e:
        print(f"\n{RED}[FAIL] Verification failed: {e}{RESET}")
        all_passed = False
        import traceback
        traceback.print_exc()
    finally:
        # Clean up temporary test uploads
        import shutil
        if test_upload_dir.exists():
            shutil.rmtree(test_upload_dir, ignore_errors=True)

    print(f"\n{BOLD}================================================================{RESET}")
    if all_passed:
        print(f"{BOLD}{GREEN}[SUCCESS] Phase 1 Ingestion Pipeline: ALL ACCEPTANCE CRITERIA MET!{RESET}")
        print(f"{BOLD}Ready for commit and Phase 2 (Understanding & Chunking).{RESET}")
        print(f"{BOLD}================================================================\n{RESET}")
        sys.exit(0)
    else:
        sys.exit(1)


if __name__ == "__main__":
    asyncio.run(run_verification())
